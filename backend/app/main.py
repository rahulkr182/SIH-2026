from dotenv import load_dotenv
load_dotenv()

import os
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from typing import Optional, List
from pydantic import BaseModel

from app.schemas import TaskResponse
from app.router.classifier import classify_task
from app.agents.graph import app_graph

app = FastAPI(title="Sovereign AI Workbench - Phase 1")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ChatMessage(BaseModel):
    role: str
    content: str

class MemoRequest(BaseModel):
    history: List[ChatMessage]

class PDFRequest(BaseModel):
    text: str

@app.post("/api/download_memo")
async def download_memo(req: MemoRequest):
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("MRPL Official Approval Memo", 0)
        
        for msg in req.history:
            p = doc.add_paragraph()
            p.add_run(f"[{msg.role.upper()}]: ").bold = True
            p.add_run(msg.content)
        
        os.makedirs("data/outputs", exist_ok=True)
        filepath = "data/outputs/memo.docx"
        doc.save(filepath)
        return FileResponse(filepath, filename="MRPL_Memo.docx")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/download_pdf")
async def download_pdf(req: PDFRequest):
    try:
        # If the agent natively edited a PDF during this session, return it directly
        filepath = "data/outputs/edited_document.pdf"
        if os.path.exists(filepath):
            return FileResponse(filepath, filename="Edited_Document.pdf")
            
        from fpdf import FPDF
        import re
        
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        
        # Clean markdown
        clean_text = req.text.replace('**', '').replace('### ', '').replace('## ', '').replace('# ', '')
        
        # Replace long dash sequences that break FPDF
        clean_text = re.sub(r'-{4,}', '---', clean_text)
        
        lines = clean_text.split('\n')
        for line in lines:
            line = line.encode('latin-1', 'replace').decode('latin-1')
            if not line.strip():
                pdf.ln(6)
                continue
            try:
                pdf.multi_cell(w=0, h=6, txt=line)
            except Exception:
                # Fallback if a line causes "not enough horizontal space"
                pdf.write(h=6, txt=line[:80])
                pdf.ln(6)
            
        os.makedirs("data/outputs", exist_ok=True)
        filepath = "data/outputs/edited_document.pdf"
        pdf.output(filepath)
        return FileResponse(filepath, filename="Edited_Document.pdf")
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/task", response_model=TaskResponse)
async def process_task(
    modality: str = Form(...),
    prompt: str = Form(...),
    file: Optional[UploadFile] = File(None)
):
    try:
        # Cleanup any previous native edits to avoid returning old files
        if os.path.exists("data/outputs/edited_document.pdf"):
            os.remove("data/outputs/edited_document.pdf")
            
        file_path = None
        if file:
            os.makedirs("data/uploads", exist_ok=True)
            file_path = os.path.join("data/uploads", file.filename)
            with open(file_path, "wb") as f:
                f.write(await file.read())
        
        # 1. Routing
        pipeline = classify_task(modality=modality, prompt=prompt, file_name=file.filename if file else None)
        
        # 2. LangGraph Execution
        initial_state = {
            "modality": modality,
            "prompt": prompt,
            "file_name": file_path,
            "trace": [f"Task routed to: {pipeline}"]
        }
        
        # We invoke the LangGraph synchronously for the hackathon prototype
        final_state = app_graph.invoke(initial_state)
        
        # Format the output trace and draft
        trace_str = "\n".join([f"- {t}" for t in final_state.get('trace', [])])
        draft = final_state.get('draft', 'No draft generated.')
        
        final_message = f"**Agent Execution Trace:**\n{trace_str}\n\n---\n\n**Final Output:**\n\n{draft}"
        
        return TaskResponse(
            status="success",
            pipeline_route=pipeline,
            message=final_message
        )
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app.main:app", host="0.0.0.0", port=8000, reload=True)
